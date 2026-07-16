"""Import a HubSpot engagement export (.xlsx) into the local store.

Reads the exact shape of the uploaded HubSpot engagement dump (emails, calls,
meetings with a ``body_preview``), runs each engagement through extraction
(offline heuristics by default, or the Claude extractor with ``use_llm=True``),
normalizes fund mentions, and persists aggregated per-contact signals — exactly
like the live ``sync`` path, but from a file instead of the HubSpot API. Still
read-only: it reads the spreadsheet, nothing writes back to HubSpot.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime

from .config import Config
from .normalize import normalize_funds
from .store import Store

MOONFARE = "moonfare.com"
NAME_EMAIL_RE = re.compile(r"([^(),;]+?)\s*\(([^)]+@[^)]+)\)")
EMAIL_TYPES = {"EMAIL", "INCOMING_EMAIL"}
CALL_TYPES = {"CALL", "MEETING"}


@dataclass
class ImportReport:
    rows: int = 0
    imported: int = 0
    skipped_type: int = 0
    skipped_no_body: int = 0
    skipped_no_contact: int = 0
    signals_written: int = 0
    servicing_skipped: int = 0
    errors: list[str] = field(default_factory=list)

    def summary(self) -> str:
        return (
            f"rows={self.rows} imported={self.imported} signals={self.signals_written} "
            f"servicing={self.servicing_skipped} skipped(type={self.skipped_type}, "
            f"no_body={self.skipped_no_body}, no_contact={self.skipped_no_contact}) "
            f"errors={len(self.errors)}"
        )


def _first_external(field_value: str | None) -> tuple[str | None, str | None]:
    """Return (name, email) of the first non-Moonfare party in an address field."""
    if not field_value:
        return None, None
    for name, email in NAME_EMAIL_RE.findall(field_value):
        email = email.strip().lower()
        if MOONFARE not in email:
            return name.strip() or None, email
    # No "Name (email)" token — try a bare email.
    for token in re.split(r"[;,]", field_value):
        token = token.strip().lower()
        if "@" in token and MOONFARE not in token:
            return None, token
    return None, None


def _to_iso(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    s = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).isoformat()
        except ValueError:
            continue
    return s.replace(" ", "T")


def import_xlsx(
    path: str,
    config: Config,
    *,
    use_llm: bool = False,
    reset: bool = True,
) -> ImportReport:
    import openpyxl

    if reset:
        import os

        if os.path.exists(config.store_path):
            os.remove(config.store_path)

    store = Store(config.store_path)
    report = ImportReport()

    extractor = None
    heuristic = None
    if use_llm:
        from .extract import Extractor

        extractor = Extractor(model=config.model)
    else:
        from .heuristics import heuristic_extract

        heuristic = heuristic_extract

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return report
    header = list(rows[0])
    # Index by first occurrence (the export has duplicate column names).
    idx: dict[str, int] = {}
    for i, name in enumerate(header):
        if name and name not in idx:
            idx[name] = i

    def cell(row, col):
        i = idx.get(col)
        return row[i] if i is not None and i < len(row) else None

    for row in rows[1:]:
        report.rows += 1
        eng_type = str(cell(row, "engagement_type") or "").upper()
        if eng_type in EMAIL_TYPES:
            object_type = "emails"
        elif eng_type in CALL_TYPES:
            object_type = "calls"
        else:
            report.skipped_type += 1
            continue

        body = str(cell(row, "body_preview") or "").strip()
        if len(body) < config.min_body_chars:
            report.skipped_no_body += 1
            continue

        # Contact = the external party. For an inbound email that's the sender;
        # otherwise the first external recipient.
        if eng_type == "INCOMING_EMAIL":
            name, email = _first_external(str(cell(row, "email_from") or ""))
        else:
            name, email = _first_external(str(cell(row, "email_to") or ""))
        if not email:
            # Fall back to whichever address field has an external party.
            for f in ("email_from", "email_to"):
                name, email = _first_external(str(cell(row, f) or ""))
                if email:
                    break
        if not email:
            report.skipped_no_contact += 1
            continue

        contact_id = email  # stable, human-readable key for the MVP
        engagement_id = str(cell(row, "engagement_id") or f"{contact_id}-{report.rows}")
        if store.is_processed(object_type, engagement_id):
            continue

        owner_first = cell(row, "first_name")
        owner_last = cell(row, "last_name")
        owner_email = cell(row, "email")
        owner_id = str(cell(row, "owner_id") or owner_email or "unknown")
        owner_name = " ".join(p for p in [owner_first, owner_last] if p).strip() or owner_email

        call_date = _to_iso(cell(row, "creation_time"))
        title = str(cell(row, "email_subject") or cell(row, "activity_type") or "").strip()

        try:
            if use_llm:
                extraction = extractor.extract(  # type: ignore[union-attr]
                    contact_id=contact_id, call_id=engagement_id,
                    call_date=call_date, title=title, body=body,
                )
            else:
                extraction = heuristic(  # type: ignore[misc]
                    contact_id=contact_id, call_id=engagement_id,
                    call_date=call_date, title=title, body=body,
                )
        except Exception as exc:  # noqa: BLE001
            report.errors.append(f"{engagement_id}: {exc}")
            continue

        report.imported += 1
        store.upsert_contact_directory(
            contact_id, name=name, email=email, lifecycle_stage=None, owner_id=owner_id
        )
        if owner_id:
            store.upsert_owner(owner_id, owner_name, str(owner_email) if owner_email else None)

        has_signal = extraction.get("funds_mentioned") or extraction.get("asset_classes")
        if not has_signal:
            report.servicing_skipped += 1
            store.mark_processed(object_type, engagement_id, contact_id, owner_id)
            continue

        normalized = normalize_funds(
            extraction.get("funds_mentioned", []), threshold=config.match_threshold
        )
        store.upsert_signal(
            contact_id=contact_id,
            contact_owner_id=owner_id,
            normalized_funds=normalized,
            extraction=extraction,
        )
        store.mark_processed(object_type, engagement_id, contact_id, owner_id)
        report.signals_written += 1

    store.close()
    return report
